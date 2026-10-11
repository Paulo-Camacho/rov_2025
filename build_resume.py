#!/usr/bin/env python3
"""Generate Paulo Camacho big-tech resume (.docx)."""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_narrow_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)


def add_hyperlink(paragraph, text, url):
    """Add a hyperlink to a paragraph (python-docx workaround)."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "20")
    r_pr.append(sz)
    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = "Calibri"
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0


def add_entry_header(doc, left, right):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r1 = p.add_run(left)
    r1.bold = True
    r1.font.size = Pt(10.5)
    r1.font.name = "Calibri"
    r2 = p.add_run("\t" + right)
    r2.font.size = Pt(10.5)
    r2.font.name = "Calibri"


def add_subheader(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.name = "Calibri"


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.15)
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.name = "Calibri"


def main():
    out_path = "/workspace/Paulo_Camacho_Resume.docx"
    doc = Document()
    set_narrow_margins(doc)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Name
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("PAULO CAMACHO")
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.name = "Calibri"
    title.paragraph_format.space_after = Pt(2)

    # Contact
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = contact.add_run(
        "831-324-7204  |  paulocamachoengineer@gmail.com  |  Marina, CA  |  "
    )
    c.font.size = Pt(9.5)
    c.font.name = "Calibri"
    add_hyperlink(contact, "paulocamacho.dev", "https://paulocamacho.dev/")
    contact.add_run("  |  ")
    add_hyperlink(contact, "LinkedIn", "https://www.linkedin.com/in/paulo-camacho-541670338")
    contact.add_run("  |  ")
    add_hyperlink(contact, "GitHub", "https://github.com/Paulo-Camacho")
    contact.paragraph_format.space_after = Pt(4)

    # Highlights strip
    hl = doc.add_paragraph()
    hl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hl.add_run(
        "Merged OSS on 17k+★ Maxun  ·  ICPC PNW Regional Div. 2 (7th)  ·  "
        "Grand Prize, Kelp Your Neighbor Hackathon"
    )
    hr.font.size = Pt(9)
    hr.font.name = "Calibri"
    hr.italic = True
    hl.paragraph_format.space_after = Pt(4)

    # EDUCATION
    add_section_heading(doc, "Education")
    add_entry_header(
        doc,
        "California State University, Monterey Bay — B.S. Computer Science",
        "Expected May 2027",
    )
    add_bullet(doc, "Dean's List")
    add_entry_header(doc, "Cuesta College — A.S. Computer Science", "May 2025")
    add_bullet(doc, "MESA STEM scholarship recipient")

    # EXPERIENCE
    add_section_heading(doc, "Experience")

    add_entry_header(
        doc,
        "Maxun (Open Source) — TypeScript, Playwright, Node.js, PostgreSQL, MinIO",
        "Jul 2026 – Aug 2026",
    )
    add_subheader(
        doc,
        "Open-Source Contributor  ·  CTI Accelerate / Summer of Open Source Engineering  ·  "
        "github.com/getmaxun/maxun (17.8k★)  ·  PR #1135 merged",
    )
    add_bullet(
        doc,
        "Reduced screenshot-enabled crawl/search runtime ~50% by eliminating a second full-browser "
        "pass (2N→N Playwright navigations); fix merged upstream after maintainer and automated review.",
    )
    add_bullet(
        doc,
        "Captured screenshots during the interpreter’s primary scrape pass and removed post-run navigation "
        "from the output post-processor so double-load regressions are structurally blocked.",
    )
    add_bullet(
        doc,
        "Validated multi-page crawls in Docker Compose (Node.js, PostgreSQL, MinIO); confirmed each URL "
        "visited once with markdown + screenshot outputs intact.",
    )

    add_entry_header(
        doc,
        "Study Buddy — TypeScript, SvelteKit, Supabase, Web Bluetooth",
        "Apr 2026",
    )
    sb_sub = doc.add_paragraph()
    sb_sub.paragraph_format.space_after = Pt(2)
    sb_sub.paragraph_format.line_spacing = 1.0
    r = sb_sub.add_run(
        "Hackathon Engineer (Team)  ·  Kelp Your Neighbor @ CSUMB  ·  "
    )
    r.italic = True
    r.font.size = Pt(10)
    r.font.name = "Calibri"
    add_hyperlink(sb_sub, "study-buddy.dev", "https://study-buddy.dev/")
    r2 = sb_sub.add_run("  ·  ")
    r2.italic = True
    r2.font.size = Pt(10)
    add_hyperlink(
        sb_sub,
        "github.com/hack-deltoideus/deltoideus",
        "https://github.com/hack-deltoideus/deltoideus",
    )
    add_bullet(
        doc,
        "Co-built Grand Prize ($1,000) biometric study companion: streams Polar H9/H10 HR/ECG over "
        "Web Bluetooth, computes stress from RMSSD and signal quality, and nudges recovery when elevated.",
    )
    add_bullet(
        doc,
        "Delivered real-time sensor/waveform UX and Supabase-backed sessions in a 48-hour SvelteKit sprint "
        "(hackathon repo); teammates later extended and hosted the public demo at study-buddy.dev.",
    )

    add_entry_header(
        doc,
        "Foundation for California Community Colleges (FoundationCCC)",
        "Jan 2025 – May 2026",
    )
    add_subheader(doc, "Research & Outreach Intern  ·  Remote, California")
    add_bullet(
        doc,
        "Built and maintained Python/SQL data pipelines ingesting and standardizing institutional datasets "
        "from four community colleges into a unified statewide analysis workflow.",
    )
    add_bullet(
        doc,
        "Cleaned and reconciled heterogeneous outreach/enrollment records to support cross-campus reporting "
        "for Cradle-to-Career research initiatives.",
    )
    add_bullet(
        doc,
        "Delivered board-facing summaries translating pipeline outputs into actionable trends for "
        "FoundationCCC leadership.",
    )

    add_entry_header(
        doc,
        "Cuesta College Robotics (Aqua Cats) — Python, PyQt5, OpenCV, Arduino",
        "Jan 2024 – Jul 2025",
    )
    add_subheader(doc, "Robotics Software Engineer  ·  San Luis Obispo, CA  ·  MATE ROV Competition")
    add_bullet(
        doc,
        "Architected Python + Arduino thruster control for competition ROV Sea Pounce: 4-motor PWM and dual "
        "claw actuation over null-terminated JSON serial (pyserial, Arduino Mega).",
    )
    add_bullet(
        doc,
        "Shipped PyQt5 ground-control station with threaded joystick (pygame), bidirectional serial I/O, "
        "and OpenCV live video for low-latency pilot feedback and mission diagnostics.",
    )
    add_bullet(
        doc,
        "Earned funded travel to represent Cuesta College at the MATE ROV regional competition.",
    )

    add_entry_header(
        doc,
        "California State University, Monterey Bay",
        "Aug 2026 – Present",
    )
    add_subheader(doc, "Teaching Assistant, CST 205 (Intro Python)  ·  Seaside, CA")
    add_bullet(
        doc,
        "Lead weekly labs on Python fundamentals and debugging; hold office hours and 1:1 sessions for "
        "students struggling with core CS concepts.",
    )
    add_bullet(
        doc,
        "Use lab performance to identify gaps early and adjust examples/explanations—improving student "
        "confidence and completion on hands-on exercises.",
    )

    # PROJECTS / COMPETITIONS (compact)
    add_section_heading(doc, "Selected Awards & Competition")
    add_bullet(doc, "ICPC Pacific Northwest Regional, Division 2 — 7th place (2025)")
    add_bullet(doc, "Kelp Your Neighbor Hackathon — Grand Prize / 1st place, Study Buddy (2026)")
    add_bullet(doc, "MESA STEM Scholarship, Cuesta College (2025)")

    # SKILLS
    add_section_heading(doc, "Technical Skills")
    skills = doc.add_paragraph()
    skills.paragraph_format.space_after = Pt(0)
    skills.paragraph_format.line_spacing = 1.05
    lines = [
        ("Languages: ", "Python, TypeScript/JavaScript, C++, Java, SQL, Bash"),
        ("Frameworks & libraries: ", "Playwright, PyQt5, SvelteKit, OpenCV, pygame, Supabase"),
        ("Systems & tools: ", "Docker, Docker Compose, Git/GitHub, Linux, PostgreSQL, MinIO, Arduino, GDB"),
    ]
    for label, rest in lines:
        r0 = skills.add_run(label)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.name = "Calibri"
        r1 = skills.add_run(rest + "\n")
        r1.font.size = Pt(10)
        r1.font.name = "Calibri"

    doc.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main()
